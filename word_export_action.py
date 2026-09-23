"""
title: Export to Word
author: rag-stack
version: 1.0.0
description: Adds an "Export answer to Word" button. Converts the Markdown answer with pandoc (headings, lists, tables, code) into a styled .docx – fully offline.
required_open_webui_version: 0.6.0
"""

# No `requirements:` line on purpose – pandoc, pypandoc and python-docx already ship in
# the standard (non-slim) Open WebUI image, so nothing is pip-installed (airgap-safe).

import asyncio
import io
import os
import re
import tempfile
import uuid
from datetime import datetime

from pydantic import BaseModel, Field

# --------------------------------------------------------------------------- #
# Pure conversion helpers (no Open WebUI imports – testable on their own)
# --------------------------------------------------------------------------- #

_DETAILS_RE = re.compile(r"<details\b[^>]*>.*?</details>", re.S | re.I)   # reasoning / tool-call blocks
_THINK_RE = re.compile(r"<think>.*?</think>", re.S | re.I)
_FENCE_RE = re.compile(r"(```.*?```|~~~.*?~~~)", re.S)                  # protect code blocks


def clean_markdown(text: str) -> str:
    """Remove Open WebUI/Qwen artefacts that should not end up in a Word file."""
    text = text or ""
    text = _DETAILS_RE.sub("", text)
    text = _THINK_RE.sub("", text)
    # Pandoc needs a blank line before a table / list that directly follows a paragraph.
    parts = _FENCE_RE.split(text)
    for i in range(0, len(parts), 2):  # only outside code fences
        parts[i] = re.sub(r"([^\n])\n(\|[^\n]*\|\s*\n\|\s*:?-{2,})", r"\1\n\n\2", parts[i])
        parts[i] = re.sub(r"([^\n|])\n((?:[-*+]|\d+\.) )", r"\1\n\n\2", parts[i])
    return "".join(parts).strip()


def sources_markdown(message: dict) -> str:
    """Numbered source list from Open WebUI citations (if the message has any)."""
    names = []
    for src in message.get("sources") or []:
        metas = src.get("metadata") or [{}]
        for meta in metas:
            name = (meta or {}).get("name") or (meta or {}).get("source") or (src.get("source") or {}).get("name")
            page = (meta or {}).get("page")
            if name:
                label = f"{name} (p. {int(page) + 1})" if isinstance(page, (int, float)) else str(name)
                if label not in names:
                    names.append(label)
    if not names:
        return ""
    return "\n\n## Sources\n\n" + "\n".join(f"{i}. {n}" for i, n in enumerate(names, 1))


def _set_cell_borders(table, size: int = 4, color: str = "999999"):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    for old in tbl_pr.findall(qn("w:tblBorders")):
        tbl_pr.remove(old)
    tbl_pr.append(borders)


def _shade(cell, fill: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _repeat_header(row):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def polish_docx(path: str, font_name: str, font_size: float, table_borders: bool, has_template: bool):
    """Make pandoc's plain output look like a business document."""
    from docx import Document
    from docx.shared import Pt

    doc = Document(path)
    if not has_template:  # a reference template defines its own fonts
        normal = doc.styles["Normal"]
        normal.font.name = font_name
        normal.font.size = Pt(font_size)
        for name in ("Body Text", "First Paragraph", "Compact"):
            if name in [s.name for s in doc.styles]:
                doc.styles[name].font.name = font_name
                doc.styles[name].font.size = Pt(font_size)
    if "Source Code" in [s.name for s in doc.styles]:  # grey background for code blocks
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        p_pr = doc.styles["Source Code"].element.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "F2F2F2")
        p_pr.append(shd)
    for table in doc.tables:
        if table_borders:
            _set_cell_borders(table)
        if table.rows:
            header = table.rows[0]
            _repeat_header(header)
            for cell in header.cells:
                _shade(cell, "E7E6E6")
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.bold = True
    doc.save(path)


def markdown_to_docx_bytes(
    markdown: str,
    title: str = "",
    reference_docx: str = "",
    font_name: str = "Calibri",
    font_size: float = 11,
    table_borders: bool = True,
    toc: bool = False,
) -> bytes:
    import pypandoc

    has_template = bool(reference_docx) and os.path.isfile(reference_docx)
    extra = ["--wrap=none"]
    if has_template:
        extra.append(f"--reference-doc={reference_docx}")
    if toc:
        extra.append("--toc")
    if title:
        extra.append(f"--metadata=title:{title}")

    with tempfile.TemporaryDirectory() as tmp:
        out = os.path.join(tmp, "out.docx")
        # 'markdown' (pandoc's own) handles pipe tables, lists, code, math.
        # yaml_metadata_block off: a model answer starting with '---' must not be parsed as YAML.
        pypandoc.convert_text(
            markdown,
            "docx",
            format="markdown-yaml_metadata_block+pipe_tables+grid_tables+strikeout+task_lists",
            outputfile=out,
            extra_args=extra,
        )
        polish_docx(out, font_name, font_size, table_borders, has_template)
        with open(out, "rb") as f:
            return f.read()


def safe_filename(text: str, fallback: str = "answer") -> str:
    base = re.sub(r"[^\w\- ]+", "", (text or "")).strip()[:50].strip().replace(" ", "_")
    stamp = datetime.now().strftime("%Y-%m-%d_%H%M")
    return f"{base or fallback}_{stamp}.docx"


def build_markdown(messages: list, message_id: str, scope: str, include_question: bool, include_sources: bool):
    """Return (markdown, title) for either the clicked answer or the whole chat."""
    if not messages:
        return "", ""
    idx = next((i for i, m in enumerate(messages) if m.get("id") == message_id), len(messages) - 1)

    if scope == "chat":
        parts, title = [], ""
        for m in messages[: idx + 1]:
            content = clean_markdown(m.get("content", ""))
            if not content:
                continue
            if m.get("role") == "user":
                title = title or content.splitlines()[0][:80]
                parts.append(f"## Question\n\n{content}")
            elif m.get("role") == "assistant":
                block = f"## Answer\n\n{content}"
                if include_sources:
                    block += sources_markdown(m).replace("## Sources", "### Sources")
                parts.append(block)
        return "\n\n".join(parts), title

    answer = messages[idx]
    question = next((m for m in reversed(messages[:idx]) if m.get("role") == "user"), None)
    title = clean_markdown(question.get("content", "")).splitlines()[0][:80] if question else ""
    md = ""
    q_text = clean_markdown(question.get("content", "")) if question else ""
    # Skip the question block when the title already shows it in full.
    if include_question and q_text and (q_text != title):
        md += f"**Question:** {q_text}\n\n---\n\n"
    md += clean_markdown(answer.get("content", ""))
    if include_sources:
        md += sources_markdown(answer)
    return md, title


# --------------------------------------------------------------------------- #
# Open WebUI Action
# --------------------------------------------------------------------------- #

# Toolbar icons (16 px, monochrome – Open WebUI inverts SVG icons in dark mode).
# Document + arrow = "export answer".
ICON_ANSWER = "data:image/svg+xml;base64,PHN2ZyB4bWxucz0iaHR0cDovL3d3dy53My5vcmcvMjAwMC9zdmciIHZpZXdCb3g9IjAgMCAyNCAyNCIgZmlsbD0ibm9uZSIgc3Ryb2tlPSIjMWYyOTM3IiBzdHJva2Utd2lkdGg9IjEuNyIgc3Ryb2tlLWxpbmVjYXA9InJvdW5kIiBzdHJva2UtbGluZWpvaW49InJvdW5kIj48cGF0aCBkPSJNMTQgM0g3YTIgMiAwIDAgMC0yIDJ2MTRhMiAyIDAgMCAwIDIgMmgxMGEyIDIgMCAwIDAgMi0yVjh6Ii8+PHBhdGggZD0iTTE0IDN2NWg1Ii8+PHBhdGggZD0iTTEyIDExdjYiLz48cGF0aCBkPSJtOSAxNCAzIDMgMy0zIi8+PC9zdmc+"

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class Action:
    class Valves(BaseModel):
        REFERENCE_DOCX: str = Field(
            default="/app/backend/data/templates/reference.docx",
            description="Optional Word template (styles, fonts, logo header). Ignored if the file does not exist.",
        )
        FONT_NAME: str = Field(default="Calibri", description="Body font when no template is used.")
        FONT_SIZE_PT: float = Field(default=11, description="Body font size when no template is used.")
        TABLE_BORDERS: bool = Field(default=True, description="Draw grid lines on all tables.")
        INCLUDE_QUESTION: bool = Field(default=True, description="Put the user's question above the answer.")
        INCLUDE_SOURCES: bool = Field(default=True, description="Append the RAG citations as a numbered list.")
        TABLE_OF_CONTENTS: bool = Field(default=False, description="Insert a table of contents (Word updates it on open).")
        TITLE_FROM_QUESTION: bool = Field(default=True, description="Use the question as document title.")

    def __init__(self):
        self.valves = self.Valves()
        # One button under every assistant message
        self.actions = [
            {"id": "answer", "name": "Export answer to Word", "icon_url": ICON_ANSWER},
        ]

    async def action(self, body: dict, __user__=None, __event_emitter__=None, __id__=None, __request__=None):
        scope = "chat" if (__id__ or "").endswith("chat") else "answer"

        async def status(text, done=False):
            if __event_emitter__:
                await __event_emitter__({"type": "status", "data": {"description": text, "done": done}})

        await status("Creating Word document…")
        try:
            md, title = build_markdown(
                body.get("messages") or [],
                body.get("id"),
                scope,
                self.valves.INCLUDE_QUESTION,
                self.valves.INCLUDE_SOURCES,
            )
            if not md.strip():
                await status("Nothing to export", done=True)
                return

            data = await asyncio.to_thread(
                markdown_to_docx_bytes,
                md,
                title if self.valves.TITLE_FROM_QUESTION else "",
                self.valves.REFERENCE_DOCX,
                self.valves.FONT_NAME,
                self.valves.FONT_SIZE_PT,
                self.valves.TABLE_BORDERS,
                self.valves.TABLE_OF_CONTENTS,
            )

            # Store like a normal upload so it is downloadable and permission-checked.
            from open_webui.models.files import FileForm, Files
            from open_webui.storage.provider import Storage

            user = __user__ or {}
            file_id = str(uuid.uuid4())
            name = safe_filename(title, "chat" if scope == "chat" else "answer")
            tags = {
                "OpenWebUI-User-Email": user.get("email", ""),
                "OpenWebUI-User-Id": user.get("id", ""),
                "OpenWebUI-User-Name": user.get("name", ""),
                "OpenWebUI-File-Id": file_id,
            }
            _, path = await asyncio.to_thread(Storage.upload_file, io.BytesIO(data), f"{file_id}_{name}", tags)
            await Files.insert_new_file(
                user.get("id"),
                FileForm(
                    id=file_id,
                    filename=name,
                    path=path,
                    data={},
                    meta={"name": name, "content_type": DOCX_MIME, "size": len(data)},
                ),
            )

            if __event_emitter__:
                # File chip under the message (saved with the chat). Click = preview + download link.
                await __event_emitter__(
                    {
                        "type": "files",
                        "data": {"files": [{"type": "file", "id": file_id, "url": file_id, "name": name, "size": len(data)}]},
                    }
                )
                await __event_emitter__(
                    {"type": "notification", "data": {"type": "success", "content": f"Word file ready: {name}"}}
                )
            await status(f"Word file ready: {name}", done=True)
        except Exception as e:
            await status(f"Word export failed: {e}", done=True)
            if __event_emitter__:
                await __event_emitter__(
                    {"type": "notification", "data": {"type": "error", "content": f"Word export failed: {e}"}}
                )
